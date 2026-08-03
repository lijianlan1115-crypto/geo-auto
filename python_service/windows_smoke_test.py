"""Windows/Mac 共用的交付前冒烟测试，不访问任何真实 AI 平台。"""

import argparse
import base64
import io
import json
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw


def build_input(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "问题"
    sheet.append(["id", "问题", "关键词", "豆包", "千问", "deepseek", "元宝", "文心一言"])
    sheet.append([1, "贵州本地干香型辣子鸡有哪些具体店铺？", "老街杨家辣子鸡", None, None, None, None, None])
    workbook.save(path)
    workbook.close()


def screenshot_data_url():
    image = Image.new("RGB", (1280, 800), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((180, 160, 1100, 650), outline=(210, 215, 225), width=3)
    draw.text((230, 260), "GEO Windows smoke-test screenshot", fill=(20, 30, 40))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")


def assert_tesseract_available():
    import pytesseract
    from ocr_checker import configure_tesseract_runtime

    command = configure_tesseract_runtime(pytesseract)
    if not command:
        raise AssertionError("Tesseract executable was not found")
    languages = set(pytesseract.get_languages(config=""))
    if "chi_sim" not in languages or "eng" not in languages:
        raise AssertionError(f"Tesseract language data incomplete: {sorted(languages)}")
    return {"command": command, "languages": sorted(languages)}


def run(expect_tesseract=False):
    with tempfile.TemporaryDirectory(prefix="geo-windows-smoke-") as temp_dir:
        root = Path(temp_dir)
        input_path = root / "input.xlsx"
        output_dir = root / "output"
        result_path = output_dir / "result.xlsx"
        database_path = output_dir / "progress.sqlite"
        build_input(input_path)

        os.environ["GEO_PROJECT_DIR"] = str(root)
        os.environ["GEO_INPUT_EXCEL"] = str(input_path)
        os.environ["GEO_OUTPUT_DIR"] = str(output_dir)
        os.environ["GEO_RESULT_EXCEL"] = str(result_path)
        os.environ["GEO_DB_PATH"] = str(database_path)
        os.environ["GEO_AI_JUDGE_ENABLED"] = "0"

        import server

        server.init_db()
        server.prepare_workbook(clear_outputs=True)
        server.seed_tasks(clear_outputs=True)

        payload = {
            "task_id": "2:qianwen",
            "row_number": 2,
            "platform": "qianwen",
            "question": "贵州本地干香型辣子鸡有哪些具体店铺？",
            "keywords": ["老街杨家辣子鸡"],
            "matched": False,
            "matched_keywords": [],
            "followup_count": 3,
            "answer_text": "冒烟测试回答，不访问真实平台。",
            "screenshot_data_url": screenshot_data_url(),
            "dom_location": {},
            "run_debug": [{"type": "windows_smoke_test"}],
        }
        saved = server.submit_result(payload)
        server.validate_workbook_file(result_path)
        embedded_before = server.embedded_cell_images(result_path)
        if len(embedded_before) != 1:
            raise AssertionError(f"expected 1 embedded image, got {len(embedded_before)}")

        screenshot_path = Path(saved["screenshot_path"])
        if not screenshot_path.is_file():
            raise AssertionError("screenshot was not written")
        screenshot_path.unlink()

        # 外部 PNG 删除后，再同步一次仍必须依靠工作簿内图片成功重建。
        if not server.sync_result_from_db():
            raise AssertionError(server.LAST_SYNC_ERROR or "result sync failed")
        server.validate_workbook_file(result_path)
        embedded_after = server.embedded_cell_images(result_path)
        if len(embedded_after) != 1:
            raise AssertionError(f"embedded image was lost after resync: {len(embedded_after)}")

        workbook = load_workbook(result_path, read_only=False, data_only=False, keep_links=False)
        try:
            sheet = server.find_question_worksheet(workbook)
            headers = server.read_headers(sheet)
            status_value = sheet.cell(row=2, column=headers["千问_状态"]).value
            if status_value not in {"未命中", "命中"}:
                raise AssertionError(f"unexpected status value: {status_value}")
        finally:
            workbook.close()

        result = {
            "ok": True,
            "embedded_images": len(embedded_after),
            "result_size": result_path.stat().st_size,
            "external_png_deleted": not screenshot_path.exists(),
        }
        if expect_tesseract:
            result["tesseract"] = assert_tesseract_available()
        return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--expect-tesseract", action="store_true")
    args = parser.parse_args()
    print(json.dumps(run(expect_tesseract=args.expect_tesseract), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
